"""Read uploaded Excel (.xlsx) or CSV sheets into rows keyed by standard column names.

Institutes name columns in many ways ("Roll No.", "Adm No", "Father Mobile"…), so headers are matched
against aliases. The first row with at least two known headers is the header row.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
from typing import Any

HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "roll": (
        "roll",
        "roll no",
        "roll number",
        "rollno",
        "roll#",
        "student id",
        "id",
        "admission no",
        "adm no",
        "enrollment no",
        "enrolment no",
        "reg no",
        "registration no",
    ),
    "name": ("name", "student", "student name", "full name", "name of student"),
    "phone": (
        "phone",
        "mobile",
        "student phone",
        "student mobile",
        "whatsapp",
        "student whatsapp",
        "contact",
        "student number",
    ),
    "parent1": (
        "parent phone",
        "parent mobile",
        "parent whatsapp",
        "parent 1",
        "parent1",
        "parent 1 phone",
        "father phone",
        "father mobile",
        "father whatsapp",
        "guardian phone",
        "guardian mobile",
        "parent",
    ),
    "parent1_name": ("parent name", "parent 1 name", "father name", "father", "guardian name", "guardian"),
    "parent2": ("parent 2", "parent2", "parent 2 phone", "mother phone", "mother mobile", "mother whatsapp"),
    "parent2_name": ("parent 2 name", "mother name", "mother"),
    "batch": ("batch", "class", "section", "group", "batch name"),
    "status": ("status", "attendance", "present", "present/absent", "p/a", "a/p", "absent", "remarks"),
    "score": ("score", "marks", "marks obtained", "obtained", "result", "total obtained"),
    "max": ("max", "max marks", "out of", "total marks", "maximum marks", "full marks", "total"),
    "test": ("test", "test name", "exam", "exam name"),
    "subject": ("subject", "course", "lecture", "period"),
    "day": ("day", "weekday"),
    "date": ("date",),
    "start": ("start", "start time", "from", "time", "timing"),
    "end": ("end", "end time", "to"),
    "teacher": ("teacher", "faculty", "teacher name", "sir", "ma'am"),
    "room": ("room", "venue", "hall", "classroom"),
}

ABSENT_WORDS = {"a", "ab", "abs", "absent", "0", "no", "n", "false", "x", "✗"}
PRESENT_WORDS = {"p", "present", "1", "yes", "y", "true", "✓", "l", "late"}


class SheetError(Exception):
    pass


def _norm_header(value: Any) -> str:
    text = re.sub(r"[\s_]+", " ", str(value or "").strip().lower())
    return text.rstrip(".:").strip()


def _alias_map() -> dict[str, str]:
    out: dict[str, str] = {}
    for key, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            out.setdefault(alias, key)
    return out


ALIASES = _alias_map()


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0):
            return value.date().isoformat()
        return value.isoformat(timespec="minutes")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, dt.time):
        return value.strftime("%H:%M")
    return str(value).strip()


def _raw_rows(filename: str, data: bytes) -> list[list[str]]:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")) or data[:2] == b"PK":
        try:
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # corrupt or not really xlsx
            raise SheetError(
                "Couldn't read that Excel file. Save it as .xlsx or .csv and try again."
            ) from exc
        ws = wb.worksheets[0]
        rows = [[cell_text(c) for c in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    if name.endswith(".xls"):
        raise SheetError("Old .xls files aren't supported. In Excel choose Save As → .xlsx (or CSV).")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    # Pick the delimiter that appears most in the first lines (title lines above the header are common).
    head = text.splitlines()[:20]
    delimiter = max((",", ";", "\t"), key=lambda d: sum(line.count(d) for line in head))
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def read_sheet(filename: str, data: bytes, max_rows: int = 5000) -> tuple[list[dict[str, str]], list[str]]:
    """Returns (rows with standard keys, recognised column keys)."""
    if not data:
        raise SheetError("The file is empty.")
    raw = _raw_rows(filename, data)
    header_idx = None
    keys: list[str | None] = []
    for idx, row in enumerate(raw[:20]):
        mapped = [ALIASES.get(_norm_header(c)) for c in row]
        if sum(1 for m in mapped if m) >= 2:
            header_idx, keys = idx, mapped
            break
    if header_idx is None:
        raise SheetError(
            "Couldn't find the header row. Put column names in the first row, e.g. Roll No, Name, Parent Phone."
        )
    # a column name used twice keeps the first occurrence (e.g. two "Phone" columns -> phone, then ignored)
    seen: set[str] = set()
    for i, key in enumerate(keys):
        if key in seen:
            keys[i] = None
        elif key:
            seen.add(key)
    rows: list[dict[str, str]] = []
    for row in raw[header_idx + 1 :]:
        if not any(c for c in row):
            continue
        item = {key: (row[i] if i < len(row) else "") for i, key in enumerate(keys) if key}
        if any(item.values()):
            rows.append(item)
        if len(rows) > max_rows:
            raise SheetError(f"The sheet has more than {max_rows} rows; split it into smaller files.")
    return rows, sorted(seen)


def is_absent(value: str) -> bool | None:
    v = value.strip().lower()
    if v in ABSENT_WORDS:
        return True
    if v in PRESENT_WORDS:
        return False
    return None


WEEKDAY_NAMES = {
    "mon": 0,
    "monday": 0,
    "tue": 1,
    "tues": 1,
    "tuesday": 1,
    "wed": 2,
    "wednesday": 2,
    "thu": 3,
    "thur": 3,
    "thurs": 3,
    "thursday": 3,
    "fri": 4,
    "friday": 4,
    "sat": 5,
    "saturday": 5,
    "sun": 6,
    "sunday": 6,
}


def parse_weekday(value: str) -> int | None:
    return WEEKDAY_NAMES.get(value.strip().lower().rstrip("."))


def parse_clock(value: str) -> dt.time | None:
    """'8:00', '08:00', '8 AM', '2:30 pm', '14.30', '2026-10-05T08:00'."""
    v = value.strip().lower().replace(".", ":")
    if "t" in v and re.match(r"\d{4}-\d{2}-\d{2}t", v):
        v = v.split("t", 1)[1]
    match = re.fullmatch(r"(\d{1,2})(?::(\d{2}))?(?::\d{2})?\s*(am|pm)?", v)
    if not match:
        return None
    hour, minute = int(match.group(1)), int(match.group(2) or 0)
    ampm = match.group(3)
    if ampm == "pm" and hour < 12:
        hour += 12
    elif ampm == "am" and hour == 12:
        hour = 0
    if hour > 23 or minute > 59:
        return None
    return dt.time(hour, minute)


def parse_sheet_date(value: str) -> dt.date | None:
    v = value.strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%d-%m-%y"):
        try:
            return dt.datetime.strptime(v, fmt).date()
        except ValueError:
            continue
    return None
